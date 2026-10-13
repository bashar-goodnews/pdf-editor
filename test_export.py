"""docx 내보내기 확인: 한글/영문 텍스트가 변환 결과에 남아 있어야 한다."""
import os
import tempfile

import pymupdf
from docx import Document

from core import PdfDocument, line_at
from exporters import export_docx

d = PdfDocument("sample.pdf")
ln = line_at(d.page(0), pymupdf.Point(100, 146))
d.edit_line(0, ln, "수정해서 내보낸 문장")  # 저장하지 않은 변경도 반영되어야 함

out = os.path.join(tempfile.mkdtemp(), "out.docx")
export_docx(d.doc.tobytes(), out)
text = "\n".join(p.text for p in Document(out).paragraphs)
print(text)
assert "수정해서 내보낸 문장" in text
assert "Hello, this is an English line." in text
assert "이 줄을 클릭" not in text
print("DOCX OK")
